#  -*- coding: utf-8 -*-
#  Copyright (c) 2026 Ivan LUCAS.
#  Previthys, application de gestion du DUERP (Document Unique d’Évaluation des Risques Professionnels).
#  Distribué sous licence GNU GPL.

import datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import ActionPrevention, CategorieRisque, Risque, Structure, UniteTravail, VersionDuerp

User = get_user_model()


class BaseTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("creer_groupes", verbosity=0)
        from django.contrib.auth.models import Group
        cls.admin = User.objects.create_user("admin", password="x", first_name="Ada", last_name="Admin")
        cls.admin.groups.add(Group.objects.get(name="Previthys - administrateur"))
        cls.lecteur = User.objects.create_user("lecteur", password="x")
        cls.lecteur.groups.add(Group.objects.get(name="Previthys - lecture"))
        cls.categorie = CategorieRisque.objects.get(nom="Autre")
        cls.unite = UniteTravail.objects.create(nom="Services techniques", effectif=10)
        cls.risque = Risque.objects.create(unite=cls.unite, categorie=cls.categorie, danger="Chute", frequence=7, gravite=10, maitrise=10)  # cotation 700 : critique

    def connexion(self, user):
        self.client.force_login(user)


class AccesTests(BaseTest):
    def test_anonyme_redirige_vers_connexion(self):
        r = self.client.get(reverse("risques_liste"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/connexion/", r["Location"])

    def test_utilisateur_sans_droit_refuse(self):
        self.connexion(User.objects.create_user("sans", password="x"))
        self.assertEqual(self.client.get(reverse("risques_liste")).status_code, 403)

    def test_lecteur_peut_lire_mais_pas_ecrire(self):
        self.connexion(self.lecteur)
        for nom in ("dashboard", "document", "unites_liste", "risques_liste", "categories_liste", "actions_liste", "versions_liste", "export_xlsx"):
            self.assertEqual(self.client.get(reverse(nom)).status_code, 200, nom)
        self.assertEqual(self.client.get(reverse("risques_ajouter")).status_code, 403)
        self.assertEqual(self.client.post(reverse("risques_supprimer", args=[self.risque.pk])).status_code, 403)
        self.assertTrue(Risque.objects.filter(pk=self.risque.pk).exists())
        self.assertNotContains(self.client.get(reverse("risques_liste")), "Supprimer")

    def test_pages_admin_ok(self):
        self.connexion(self.admin)
        for nom in ("unites_ajouter", "categories_ajouter", "risques_ajouter", "actions_ajouter", "versions_ajouter"):
            self.assertEqual(self.client.get(reverse(nom)).status_code, 200, nom)


class CrudTests(BaseTest):
    def setUp(self):
        self.connexion(self.admin)

    def test_niveau_et_cotation(self):
        self.assertEqual((self.risque.cotation, self.risque.niveau), (700, "critique"))
        # maîtrise non précisée : valeur par défaut du modèle (10, non maîtrisé)
        self.assertEqual(Risque(gravite=1, frequence=1).cotation, 10)
        self.assertEqual(Risque(gravite=1, frequence=1).niveau, "faible")
        self.assertEqual(Risque(gravite=4, frequence=4).cotation, 160)
        self.assertEqual(Risque(gravite=4, frequence=4).niveau, "moyen")
        self.assertEqual(Risque(gravite=10, frequence=10, maitrise=10).niveau, "critique")

    def test_ajout_risque_et_bouton_ajouter(self):
        r = self.client.post(reverse("risques_ajouter"), {"unite": self.unite.pk, "categorie": self.categorie.pk, "danger": "Bruit", "frequence": 4, "gravite": 4, "maitrise": 4, "enregistrer_ajouter": "1"})
        self.assertRedirects(r, reverse("risques_ajouter"))
        self.assertTrue(Risque.objects.filter(danger="Bruit").exists())

    def test_categorie_doublon_refusee(self):
        r = self.client.post(reverse("categories_ajouter"), {"nom": "autre", "ordre": 1})
        self.assertContains(r, "Cette catégorie existe déjà.")
        self.assertEqual(CategorieRisque.objects.filter(nom__iexact="autre").count(), 1)
        r = self.client.post(reverse("categories_ajouter"), {"nom": "Travail isolé", "ordre": 12})
        self.assertRedirects(r, reverse("categories_liste"))
        modif = CategorieRisque.objects.get(nom="Travail isolé")
        r = self.client.post(reverse("categories_modifier", args=[modif.pk]), {"nom": "Travail isolé", "ordre": 13})
        self.assertRedirects(r, reverse("categories_liste"))  # se renommer soi-même reste possible

    def test_categorie_description_facultative_et_affichee(self):
        r = self.client.post(reverse("categories_ajouter"), {"nom": "Travail en hauteur", "description": "Interventions sur toiture, échafaudage ou nacelle.", "ordre": 5})
        self.assertRedirects(r, reverse("categories_liste"))
        cat = CategorieRisque.objects.get(nom="Travail en hauteur")
        self.assertEqual(cat.description, "Interventions sur toiture, échafaudage ou nacelle.")
        self.assertContains(self.client.get(reverse("categories_liste")), "Interventions sur toiture")
        # sans description : toujours acceptée (champ facultatif)
        r = self.client.post(reverse("categories_ajouter"), {"nom": "Sans description", "ordre": 6})
        self.assertRedirects(r, reverse("categories_liste"))
        self.assertEqual(CategorieRisque.objects.get(nom="Sans description").description, "")

    def test_categorie_utilisee_non_supprimable(self):
        page = self.client.get(reverse("categories_supprimer", args=[self.categorie.pk]))
        self.assertContains(page, "Suppression impossible")
        self.client.post(reverse("categories_supprimer", args=[self.categorie.pk]), follow=True)
        self.assertTrue(CategorieRisque.objects.filter(pk=self.categorie.pk).exists())

    def test_categorie_libre_supprimable(self):
        libre = CategorieRisque.objects.create(nom="Temporaire", ordre=99)
        self.client.post(reverse("categories_supprimer", args=[libre.pk]))
        self.assertFalse(CategorieRisque.objects.filter(pk=libre.pk).exists())

    def test_suppression_unite_annonce_les_dependances(self):
        ActionPrevention.objects.create(risque=self.risque, description="Remplacer l'échelle")
        page = self.client.get(reverse("unites_supprimer", args=[self.unite.pk]))
        self.assertContains(page, "1 risque")
        self.assertContains(page, "1 action de prévention")
        self.client.post(reverse("unites_supprimer", args=[self.unite.pk]))
        self.assertFalse(Risque.objects.filter(pk=self.risque.pk).exists())
        self.assertEqual(ActionPrevention.objects.count(), 0)

    def test_action_en_retard(self):
        hier = datetime.date.today() - datetime.timedelta(days=1)
        retard = ActionPrevention.objects.create(risque=self.risque, description="A", echeance=hier)
        fait = ActionPrevention.objects.create(risque=self.risque, description="B", echeance=hier, statut="terminee")
        self.assertTrue(retard.en_retard)
        self.assertFalse(fait.en_retard)
        page = self.client.get(reverse("actions_liste"))
        self.assertContains(page, "En retard")
        self.assertContains(self.client.get(reverse("dashboard")), "En retard")

    def test_liste_risques_trie_par_cotation(self):
        Risque.objects.create(unite=self.unite, categorie=self.categorie, danger="Petit", gravite=1, frequence=1, maitrise=1)
        page = self.client.get(reverse("risques_liste")).content.decode()
        self.assertIn('data-order="700"', page)   # self.risque : 7 x 10 x 10
        self.assertIn('data-order="1"', page)     # "Petit" : 1 x 1 x 1
        self.assertIn('data-ordre="7,desc"', page)


class VersionsEtExportTests(BaseTest):
    def setUp(self):
        self.connexion(self.admin)

    def test_archivage_fige_les_donnees(self):
        r = self.client.post(reverse("versions_ajouter"), {"commentaire": "Première version"})
        self.assertRedirects(r, reverse("versions_liste"))
        version = VersionDuerp.objects.get()
        self.assertEqual((version.numero, version.auteur), (1, self.admin))
        self.assertEqual(version.donnees[0]["risques"][0]["danger"], "Chute")
        self.risque.danger = "Chute modifiée"
        self.risque.save()
        page = self.client.get(reverse("versions_document", args=[version.pk]))
        self.assertContains(page, "Chute")
        self.assertNotContains(page, "Chute modifiée")
        self.client.post(reverse("versions_ajouter"), {"commentaire": "Deuxième"})
        self.assertEqual(VersionDuerp.objects.first().numero, 2)

    def test_motif_obligatoire(self):
        r = self.client.post(reverse("versions_ajouter"), {"commentaire": ""})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(VersionDuerp.objects.count(), 0)

    def test_export_xlsx(self):
        ActionPrevention.objects.create(risque=self.risque, description="Formation", responsable=self.admin, echeance=datetime.date(2030, 1, 1))
        r = self.client.get(reverse("export_xlsx"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Synthèse", "Risques", "Actions"])
        risques, actions = wb["Risques"], wb["Actions"]
        self.assertEqual(risques["C2"].value, "Chute")
        self.assertEqual(risques["H2"].value, "=E2*F2*G2")
        self.assertEqual(actions["D2"].value, "Ada Admin")

    def test_export_version_et_version_inconnue(self):
        self.client.post(reverse("versions_ajouter"), {"commentaire": "v1"})
        v = VersionDuerp.objects.get()
        self.assertEqual(self.client.get(reverse("versions_export", args=[v.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("versions_export", args=[999])).status_code, 404)


class StructuresTests(BaseTest):
    def test_un_utilisateur_ne_voit_que_ses_structures(self):
        a, b = Structure.objects.create(nom="Mairie A"), Structure.objects.create(nom="Mairie B")
        UniteTravail.objects.create(nom="Unité de A", structure=a)
        ub = UniteTravail.objects.create(nom="Unité de B", structure=b)
        Risque.objects.create(unite=ub, categorie=self.categorie, danger="Risque de B")
        from django.contrib.auth.models import Group
        user = User.objects.create_user("redacteur", password="x")
        user.groups.add(Group.objects.get(name="Previthys - rédacteur"))
        user.structures.add(a)
        self.connexion(user)
        page = self.client.get(reverse("unites_liste"))
        self.assertContains(page, "Unité de A")
        self.assertContains(page, "Services techniques")       # sans structure : visible de tous
        self.assertNotContains(page, "Unité de B")
        self.assertNotContains(self.client.get(reverse("risques_liste")), "Risque de B")
        self.assertEqual(self.client.get(reverse("unites_modifier", args=[ub.pk])).status_code, 404)

    def test_superuser_voit_tout(self):
        b = Structure.objects.create(nom="Mairie B")
        UniteTravail.objects.create(nom="Unité de B", structure=b)
        self.connexion(User.objects.create_superuser("root", password="x"))
        self.assertContains(self.client.get(reverse("unites_liste")), "Unité de B")

    def test_archivage_limite_a_la_structure(self):
        a, b = Structure.objects.create(nom="Mairie A"), Structure.objects.create(nom="Mairie B")
        UniteTravail.objects.create(nom="Unité de A", structure=a)
        UniteTravail.objects.create(nom="Unité de B", structure=b)
        self.connexion(User.objects.create_superuser("root", password="x"))
        self.client.post(reverse("versions_ajouter"), {"structure": a.pk, "commentaire": "Archive A"})
        noms = [u["nom"] for u in VersionDuerp.objects.get().donnees]
        self.assertIn("Unité de A", noms)
        self.assertIn("Services techniques", noms)
        self.assertNotIn("Unité de B", noms)


class GroupesTests(TestCase):
    def test_commande_creer_groupes(self):
        call_command("creer_groupes", verbosity=0)
        from django.contrib.auth.models import Group
        self.assertEqual(Group.objects.filter(name__startswith="Previthys").count(), 3)
        lecture = Group.objects.get(name="Previthys - lecture")
        self.assertTrue(lecture.permissions.filter(codename="view_risque").exists())
        self.assertFalse(lecture.permissions.filter(codename="add_risque").exists())


class ExempleTests(TestCase):
    def test_charger_exemple_est_idempotent(self):
        call_command("charger_exemple", verbosity=0)
        self.assertEqual((UniteTravail.objects.count(), Risque.objects.count(), ActionPrevention.objects.count()), (3, 7, 4))
        call_command("charger_exemple", verbosity=0)
        self.assertEqual(UniteTravail.objects.count(), 3)


class MesuresRealiseesTests(BaseTest):
    """Une action terminée devient une mesure existante ; une action non terminée reste « prévue »."""

    def setUp(self):
        self.connexion(self.admin)
        self.realisee = ActionPrevention.objects.create(risque=self.risque, description="Échelles remplacées", statut="terminee", date_realisation=datetime.date(2026, 3, 4))
        self.prevue = ActionPrevention.objects.create(risque=self.risque, description="Formation gestes et postures", statut="a_faire", echeance=datetime.date(2030, 1, 1))
        self.risque.mesures_existantes = "Échelles contrôlées."
        self.risque.save()

    def test_normaliser_separe_realisees_et_prevues(self):
        from core.utils import construire_donnees, normaliser_donnees
        r = normaliser_donnees(construire_donnees(self.admin))[0]["risques"][0]
        self.assertEqual([m["description"] for m in r["mesures_realisees"]], ["Échelles remplacées"])
        self.assertEqual(r["mesures_realisees"][0]["date"], "04/03/2026")
        self.assertEqual([a["description"] for a in r["actions"]], ["Formation gestes et postures"])
        self.assertEqual(len(r["actions_toutes"]), 2)

    def test_document_affiche_la_mesure_realisee_dans_les_mesures_existantes(self):
        html = self.client.get(reverse("document")).content.decode()
        self.assertIn("Échelles contrôlées.", html)
        self.assertIn("• Échelles remplacées", html)
        self.assertIn("réalisée le 04/03/2026", html)
        # l'action terminée n'apparaît plus dans « Actions prévues » : une seule occurrence de son texte
        self.assertEqual(html.count("Échelles remplacées"), 1)
        self.assertIn("Formation gestes et postures", html)

    def test_anciennes_archives_sans_indicateur_restent_lisibles(self):
        ancienne = [{"nom": "U", "effectif": 1, "description": "", "risques": [{
            "categorie": "Autre", "danger": "Ancien danger", "situation": "", "gravite": 2, "frequence": 2, "cotation": 4, "niveau": "moyen",
            "mesures_existantes": "", "actions": [
                {"description": "Faite", "statut": "Terminée", "echeance": "", "responsable": ""},
                {"description": "À venir", "statut": "À faire", "echeance": "", "responsable": ""}]}]}]
        v = VersionDuerp.objects.create(numero=9, commentaire="ancienne", donnees=ancienne)
        html = self.client.get(reverse("versions_document", args=[v.pk])).content.decode()
        self.assertIn("• Faite", html)
        self.assertIn("(réalisée)", html)
        self.assertIn("À venir", html)
        self.assertEqual(self.client.get(reverse("versions_export", args=[v.pk])).status_code, 200)

    def test_date_de_realisation_automatique_puis_effacee(self):
        data = {"risque": self.risque.pk, "description": "X", "statut": "terminee"}
        self.client.post(reverse("actions_ajouter"), data)
        a = ActionPrevention.objects.get(description="X")
        self.assertEqual(a.date_realisation, datetime.date.today())
        data["statut"] = "en_cours"
        self.client.post(reverse("actions_modifier", args=[a.pk]), data)
        a.refresh_from_db()
        self.assertIsNone(a.date_realisation)

    def test_date_saisie_conservee(self):
        self.client.post(reverse("actions_ajouter"), {"risque": self.risque.pk, "description": "Y", "statut": "terminee", "date_realisation": "2026-01-15"})
        self.assertEqual(ActionPrevention.objects.get(description="Y").date_realisation, datetime.date(2026, 1, 15))

    def test_invitation_a_reevaluer_apres_passage_a_terminee(self):
        r = self.client.post(reverse("actions_modifier", args=[self.prevue.pk]), {"risque": self.risque.pk, "description": self.prevue.description, "statut": "terminee"}, follow=True)
        self.assertContains(r, "Réévaluer le risque")
        self.assertContains(r, reverse("risques_modifier", args=[self.risque.pk]))
        # une modification qui ne change pas le statut n'invite pas à réévaluer
        r = self.client.post(reverse("actions_modifier", args=[self.prevue.pk]), {"risque": self.risque.pk, "description": "Reformulée", "statut": "terminee"}, follow=True)
        self.assertNotContains(r, "Réévaluer le risque")

    def test_export_xlsx_mesures_et_date_de_realisation(self):
        r = self.client.get(reverse("export_xlsx"))
        wb = load_workbook(BytesIO(r.content))
        mesures = wb["Risques"]["J2"].value
        self.assertIn("Échelles contrôlées.", mesures)
        self.assertIn("• Échelles remplacées (réalisée le 04/03/2026)", mesures)
        self.assertNotIn("Formation", mesures)
        actions = wb["Actions"]
        self.assertEqual(actions["H1"].value, "Date de réalisation")
        lignes = {actions.cell(row=i, column=3).value: actions.cell(row=i, column=8).value for i in range(2, 4)}
        self.assertEqual(lignes["Échelles remplacées"].date(), datetime.date(2026, 3, 4))
        self.assertIsNone(lignes["Formation gestes et postures"])


class ConnexionTests(BaseTest):
    def test_page_de_connexion_accessible_sans_etre_connecte(self):
        r = self.client.get(reverse("login"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Se connecter")
        self.assertTemplateUsed(r, "core/login.html")

    def test_connexion_reussie_redirige_vers_le_tableau_de_bord(self):
        r = self.client.post(reverse("login"), {"username": "admin", "password": "x"})
        self.assertRedirects(r, reverse("dashboard"))
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)

    def test_connexion_revient_a_la_page_demandee(self):
        r = self.client.get(reverse("risques_liste"))
        self.assertRedirects(r, "%s?next=%s" % (reverse("login"), reverse("risques_liste")))
        r = self.client.post("%s?next=%s" % (reverse("login"), reverse("risques_liste")), {"username": "admin", "password": "x", "next": reverse("risques_liste")})
        self.assertRedirects(r, reverse("risques_liste"))

    def test_mauvais_mot_de_passe(self):
        r = self.client.post(reverse("login"), {"username": "admin", "password": "faux"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Identifiant ou mot de passe incorrect.")
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 302)

    def test_utilisateur_deja_connecte_renvoye_au_tableau_de_bord(self):
        self.connexion(self.admin)
        self.assertRedirects(self.client.get(reverse("login")), reverse("dashboard"))

    def test_deconnexion_par_post(self):
        self.connexion(self.admin)
        r = self.client.post(reverse("logout"))
        self.assertRedirects(r, reverse("login"))
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 302)

    def test_deconnexion_par_get_refusee(self):
        self.connexion(self.admin)
        self.assertEqual(self.client.get(reverse("logout")).status_code, 405)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)


class Bootstrap5Tests(TestCase):
    """Garde-fou de la migration vers Bootstrap 5 : plus de syntaxe Bootstrap 4 dans les gabarits, les formulaires ni les vues."""

    def test_aucune_syntaxe_bootstrap4(self):
        import pathlib
        import re
        racine = pathlib.Path(__file__).parent
        interdits = re.compile(r"data-toggle|data-target|data-dismiss|badge-(?:primary|secondary|success|danger|warning|info|light|dark)|custom-select|form-group|form-inline|btn-block|class=.close.|\bmr-auto\b|\bml-auto\b|text-left|text-right|sr-only|dropdown-menu-right|btn-default")
        fichiers = [f for pattern in ("templates/**/*.html", "static/core/*.js", "static/core/*.css", "forms/*.py", "views/*.py") for f in racine.glob(pattern)]
        self.assertTrue(fichiers)
        for f in fichiers:
            trouve = interdits.search(f.read_text(encoding="utf-8"))
            self.assertIsNone(trouve, "syntaxe Bootstrap 4 dans %s : %s" % (f.relative_to(racine), trouve and trouve.group(0)))

    def test_les_pages_chargent_bootstrap_5(self):
        admin = User.objects.create_superuser("root", password="x")
        self.client.force_login(admin)
        page = self.client.get(reverse("risques_liste")).content.decode()
        for fichier in ("bootstrap.min.css", "bootstrap.bundle.min.js", "dataTables.bootstrap5.min.js", "buttons.bootstrap5.min.js"):
            self.assertIn(fichier, page)
        self.assertNotIn("bootstrap4", page)
        import pathlib
        vendor = pathlib.Path(__file__).parent / "static" / "core" / "vendor"
        self.assertRegex((vendor / "bootstrap.min.css").read_text(encoding="utf-8")[:300], r"Bootstrap\s+v5\.3")


class DashboardWidgetsTests(BaseTest):
    def setUp(self):
        self.connexion(self.admin)

    def test_top_risques_les_plus_critiques(self):
        Risque.objects.create(unite=self.unite, categorie=self.categorie, danger="Petit risque", gravite=1, frequence=1, maitrise=1)
        page = self.client.get(reverse("dashboard"))
        self.assertContains(page, "Risques les plus critiques")
        self.assertContains(page, "Chute")   # self.risque : 700, doit apparaître avant "Petit risque" (1)
        contenu = page.content.decode()
        self.assertLess(contenu.index("Chute"), contenu.index("Petit risque"))
        self.assertContains(page, reverse("risques_modifier", args=[self.risque.pk]))

    def test_repartition_par_categorie(self):
        autre_categorie = CategorieRisque.objects.create(nom="Bruit", ordre=99)
        Risque.objects.create(unite=self.unite, categorie=autre_categorie, danger="Bruit machine", gravite=1, frequence=1, maitrise=1)
        page = self.client.get(reverse("dashboard"))
        self.assertContains(page, "Répartition par catégorie de risque")
        self.assertContains(page, "Bruit")
        self.assertContains(page, self.categorie.nom)

    def test_lecteur_sans_droit_de_modifier_voit_le_texte_sans_lien(self):
        from django.contrib.auth.models import Group
        lecteur = User.objects.create_user("lecteur2", password="x")
        lecteur.groups.add(Group.objects.get(name="Previthys - lecture"))
        self.client.force_login(lecteur)
        page = self.client.get(reverse("dashboard"))
        self.assertContains(page, "Chute")
        self.assertNotContains(page, reverse("risques_modifier", args=[self.risque.pk]))
