"""Tests des actions générales (sans risque) et communes (plusieurs risques / plusieurs unités)."""
#  -*- coding: utf-8 -*-
#  Copyright (c) 2026 Ivan LUCAS.
#  Previthys, application de gestion du DUERP (Document Unique d’Évaluation des Risques Professionnels).
#  Distribué sous licence GNU GPL.

import json
import tempfile
import zipfile
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import ActionPrevention, CategorieAction, CategorieRisque, JournalAudit, Risque, Structure, UniteTravail, VersionDuerp

User = get_user_model()


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("creer_groupes", verbosity=0)
        cls.struct_a = Structure.objects.create(nom="Mairie A")
        cls.struct_b = Structure.objects.create(nom="Mairie B")
        cls.root = User.objects.create_superuser("root", password="x")
        cls.user_a = User.objects.create_user("user_a", password="x")
        cls.user_a.groups.add(Group.objects.get(name="Previthys - administrateur"))
        cls.user_a.structures.add(cls.struct_a)
        cls.cat = CategorieRisque.objects.get(nom="Risques psychosociaux")
        cls.voirie = UniteTravail.objects.create(nom="Voirie", structure=cls.struct_a)
        cls.accueil = UniteTravail.objects.create(nom="Accueil", structure=cls.struct_a)
        cls.ecole_b = UniteTravail.objects.create(nom="École B", structure=cls.struct_b)
        cls.rps_voirie = Risque.objects.create(unite=cls.voirie, categorie=cls.cat, danger="Stress voirie", frequence=4, gravite=4, maitrise=4)
        cls.rps_accueil = Risque.objects.create(unite=cls.accueil, categorie=cls.cat, danger="Agressions accueil", frequence=7, gravite=4, maitrise=4)
        cls.rps_b = Risque.objects.create(unite=cls.ecole_b, categorie=cls.cat, danger="Stress école B", frequence=4, gravite=4, maitrise=4)

    def setUp(self):
        self.client.force_login(self.root)


class SaisieTests(Base):
    def test_action_commune_une_seule_action_pour_plusieurs_unites(self):
        self.client.post(reverse("actions_ajouter"), {"portee": "risques", "risques": [self.rps_voirie.pk, self.rps_accueil.pk],
                                                      "description": "Formation stress", "statut": "a_faire", "cout": "1100"})
        action = ActionPrevention.objects.get(description="Formation stress")
        self.assertEqual(set(action.risques.all()), {self.rps_voirie, self.rps_accueil})
        self.assertTrue(action.est_commune)
        self.assertEqual(action.structure, self.struct_a)   # déduite des unités
        for risque in (self.rps_voirie, self.rps_accueil):   # visible sur la fiche de chaque risque
            self.assertContains(self.client.get(reverse("risques_modifier", args=[risque.pk])), "Commune à 2 unités")
        self.assertContains(self.client.get(reverse("actions_liste")), "Commune à 2 unités")

    def test_action_generale_sans_risque(self):
        self.client.post(reverse("actions_ajouter"), {"portee": "generale", "risques": [self.rps_voirie.pk], "structure": self.struct_a.pk,
                                                      "description": "Utilisation de OpenGST", "statut": "a_faire"})
        action = ActionPrevention.objects.get(description="Utilisation de OpenGST")
        self.assertFalse(action.risques.exists())             # les cases cochées sont ignorées pour une action générale
        self.assertEqual(action.structure, self.struct_a)
        self.assertContains(self.client.get(reverse("actions_liste") + "?portee=generales"), "Utilisation de OpenGST")
        self.assertNotContains(self.client.get(reverse("actions_liste") + "?portee=communes"), "Utilisation de OpenGST")
        document = self.client.get(reverse("document"))
        self.assertContains(document, "Mesures générales de prévention")
        self.assertContains(document, "Utilisation de OpenGST")

    def test_portee_risques_sans_risque_coche_refusee(self):
        r = self.client.post(reverse("actions_ajouter"), {"portee": "risques", "description": "Sans risque", "statut": "a_faire"})
        self.assertContains(r, "Cochez au moins un risque")
        self.assertFalse(ActionPrevention.objects.filter(description="Sans risque").exists())

    def test_journal_indique_le_changement_de_risques(self):
        action = ActionPrevention.objects.create(risque=self.rps_voirie, description="Former")
        self.client.post(reverse("actions_modifier", args=[action.pk]), {"portee": "risques", "risques": [self.rps_voirie.pk, self.rps_accueil.pk],
                                                                         "description": "Former", "statut": "a_faire"})
        self.assertIn("Risques concernés : 1 → 2", JournalAudit.objects.filter(action="modification").latest("pk").detail)


class FicheRisqueTests(Base):
    def test_detacher_une_action_commune_la_conserve_pour_les_autres(self):
        action = ActionPrevention.objects.create(risques=[self.rps_voirie, self.rps_accueil], description="Commune")
        self.client.post(reverse("risques_actions_supprimer", args=[self.rps_voirie.pk, action.pk]))
        self.assertEqual(list(ActionPrevention.objects.get(pk=action.pk).risques.all()), [self.rps_accueil])
        # sur son dernier risque, le bouton supprime l'action
        self.client.post(reverse("risques_actions_supprimer", args=[self.rps_accueil.pk, action.pk]))
        self.assertFalse(ActionPrevention.objects.filter(pk=action.pk).exists())

    def test_rattacher_une_action_existante(self):
        action = ActionPrevention.objects.create(risque=self.rps_voirie, description="À partager")
        page = self.client.get(reverse("risques_modifier", args=[self.rps_accueil.pk]))
        self.assertContains(page, "Rattacher une action existante")
        self.client.post(reverse("risques_actions_rattacher", args=[self.rps_accueil.pk]), {"action": action.pk})
        self.assertEqual(action.risques.count(), 2)

    def test_ajout_depuis_la_fiche_lie_l_action_au_risque(self):
        self.client.post(reverse("risques_actions_ajouter", args=[self.rps_voirie.pk]), {"description": "Depuis la fiche", "statut": "a_faire"})
        self.assertEqual(list(ActionPrevention.objects.get(description="Depuis la fiche").risques.all()), [self.rps_voirie])


class SuppressionTests(Base):
    def test_suppression_d_un_risque(self):
        propre = ActionPrevention.objects.create(risque=self.rps_voirie, description="Propre")
        commune = ActionPrevention.objects.create(risques=[self.rps_voirie, self.rps_accueil], description="Commune")
        page = self.client.get(reverse("risques_supprimer", args=[self.rps_voirie.pk]))
        self.assertContains(page, "1 action de prévention")
        self.assertContains(page, "1 action commune détachée")
        self.client.post(reverse("risques_supprimer", args=[self.rps_voirie.pk]))
        self.assertFalse(ActionPrevention.objects.filter(pk=propre.pk).exists())
        self.assertEqual(list(ActionPrevention.objects.get(pk=commune.pk).risques.all()), [self.rps_accueil])

    def test_suppression_d_une_unite_ne_laisse_pas_d_action_orpheline(self):
        autre = Risque.objects.create(unite=self.voirie, categorie=self.cat, danger="Autre risque voirie", frequence=1, gravite=1, maitrise=1)
        action = ActionPrevention.objects.create(risques=[self.rps_voirie, autre], description="Deux risques, même unité")
        self.voirie.delete()
        self.assertFalse(ActionPrevention.objects.filter(pk=action.pk).exists())   # ne devient pas une action générale


class VisibiliteTests(Base):
    def setUp(self):
        self.client.force_login(self.user_a)

    def test_actions_generales_par_structure(self):
        ActionPrevention.objects.create(description="Générale A", structure=self.struct_a)
        ActionPrevention.objects.create(description="Générale B", structure=self.struct_b)
        ActionPrevention.objects.create(description="Générale pour tous")
        liste = self.client.get(reverse("actions_liste")).content.decode()
        self.assertIn("Générale A", liste)
        self.assertIn("Générale pour tous", liste)
        self.assertNotIn("Générale B", liste)
        self.assertNotIn("Générale B", self.client.get(reverse("document")).content.decode())

    def test_action_commune_a_deux_structures(self):
        action = ActionPrevention.objects.create(risques=[self.rps_voirie, self.rps_b], description="Commune A et B")
        liste = self.client.get(reverse("actions_liste")).content.decode()
        self.assertIn("Commune A et B", liste)
        self.assertNotIn("École B", liste)   # l'unité de l'autre structure n'est pas nommée
        # l'enregistrement par un utilisateur de A conserve le risque de B, qu'il ne voit pas
        self.client.post(reverse("actions_modifier", args=[action.pk]), {"portee": "risques", "risques": [self.rps_voirie.pk],
                                                                         "description": "Commune A et B", "statut": "en_cours"})
        self.assertEqual(set(action.risques.all()), {self.rps_voirie, self.rps_b})
        # et elle ne peut pas devenir générale
        r = self.client.post(reverse("actions_modifier", args=[action.pk]), {"portee": "generale", "description": "X", "statut": "a_faire"})
        self.assertContains(r, "ne peut pas devenir une action générale")

    def test_generale_d_une_autre_structure_inaccessible(self):
        b = ActionPrevention.objects.create(description="Générale B", structure=self.struct_b)
        self.assertEqual(self.client.get(reverse("actions_modifier", args=[b.pk])).status_code, 404)


class ArchivesEtExportTests(Base):
    def test_version_archivee_et_export(self):
        ActionPrevention.objects.create(description="Sensibilisation générale")
        ActionPrevention.objects.create(risques=[self.rps_voirie, self.rps_accueil], description="Formation commune")
        self.client.post(reverse("versions_ajouter"), {"commentaire": "Mise à jour annuelle"})
        version = VersionDuerp.objects.get()
        self.assertEqual([m["description"] for m in version.mesures_generales], ["Sensibilisation générale"])
        self.assertContains(self.client.get(reverse("versions_document", args=[version.pk])), "Sensibilisation générale")
        wb = load_workbook(BytesIO(self.client.get(reverse("export_xlsx")).content))
        valeurs = [c.value for ligne in wb["Actions"].iter_rows() for c in ligne]
        self.assertIn("Mesure générale", valeurs)
        self.assertIn("Commune à 2 unités", valeurs)

    def test_ancienne_archive_sans_mesures_generales(self):
        version = VersionDuerp.objects.create(numero=1, commentaire="Ancienne", donnees=[])
        self.assertEqual(self.client.get(reverse("versions_document", args=[version.pk])).status_code, 200)


class ImportTests(TestCase):
    def test_import_format_2(self):
        donnees = {
            "format": "previthys-import/2", "source": "test",
            "unites": [{"nom": "U1"}, {"nom": "U2"}], "categories": [{"nom": "Autre"}],
            "risques": [{"id": 1, "unite": "U1", "categorie": "Autre", "danger": "D1", "frequence": 4, "gravite": 4, "maitrise": 4},
                        {"id": 2, "unite": "U2", "categorie": "Autre", "danger": "D2", "frequence": 4, "gravite": 4, "maitrise": 4}],
            "actions": [{"description": "Commune", "risques": [1, 2], "cout": 1100}, {"description": "Générale", "risques": []},
                        {"description": "Propre", "risques": [1]}],
        }
        with tempfile.NamedTemporaryFile(suffix=".zip") as f:
            with zipfile.ZipFile(f.name, "w") as z:
                z.writestr("donnees.json", json.dumps(donnees))
            sortie = StringIO()
            call_command("importer_duerp", f.name, stdout=sortie)
        self.assertIn("dont 1 générales et 1 communes", sortie.getvalue())
        self.assertEqual(ActionPrevention.objects.get(description="Commune").risques.count(), 2)
        self.assertFalse(ActionPrevention.objects.get(description="Générale").risques.exists())


class CategoriesActionsTests(Base):
    def test_categories_par_defaut_et_droits_des_groupes(self):
        self.assertEqual(list(CategorieAction.objects.values_list("nom", flat=True)),
                         ["Formation", "Matériel", "Organisation", "Travaux", "Communication", "Autre"])
        self.assertTrue(self.user_a.has_perm("core.change_categorieaction"))
        self.client.force_login(self.user_a)
        self.assertContains(self.client.get(reverse("categories_actions_liste")), "Formation")

    def test_saisie_filtre_et_affichage(self):
        formation = CategorieAction.objects.get(nom="Formation")
        self.client.post(reverse("actions_ajouter"), {"portee": "risques", "risques": [self.rps_voirie.pk], "categorie": formation.pk,
                                                      "description": "Formation stress", "statut": "a_faire"})
        ActionPrevention.objects.create(risque=self.rps_voirie, description="Sans catégorie")
        self.assertEqual(ActionPrevention.objects.get(description="Formation stress").categorie, formation)
        filtree = self.client.get(reverse("actions_liste") + "?categorie=%d" % formation.pk).content.decode()
        self.assertIn("Formation stress", filtree)
        self.assertNotIn("Sans catégorie", filtree)
        non_classees = self.client.get(reverse("actions_liste") + "?categorie=aucune").content.decode()
        self.assertIn("Sans catégorie", non_classees)
        self.assertNotIn("Formation stress", non_classees)
        self.assertContains(self.client.get(reverse("risques_modifier", args=[self.rps_voirie.pk])), "categorie-action")
        wb = load_workbook(BytesIO(self.client.get(reverse("export_xlsx")).content))
        self.assertIn("Formation", [c.value for c in wb["Actions"]["J"]])

    def test_categorie_utilisee_non_supprimable(self):
        formation = CategorieAction.objects.get(nom="Formation")
        ActionPrevention.objects.create(risque=self.rps_voirie, description="X", categorie=formation)
        self.client.post(reverse("categories_actions_supprimer", args=[formation.pk]))
        self.assertTrue(CategorieAction.objects.filter(pk=formation.pk).exists())

    def test_import_avec_categorie(self):
        donnees = {"format": "previthys-import/2", "unites": [{"nom": "U"}], "categories": [{"nom": "Autre"}],
                   "risques": [{"id": 1, "unite": "U", "categorie": "Autre", "danger": "D", "frequence": 1, "gravite": 1, "maitrise": 1}],
                   "actions": [{"description": "A1", "risques": [1], "categorie": "Formation"}, {"description": "A2", "risques": [], "categorie": "Nouvelle"}]}
        with tempfile.NamedTemporaryFile(suffix=".zip") as f:
            with zipfile.ZipFile(f.name, "w") as z:
                z.writestr("donnees.json", json.dumps(donnees))
            call_command("importer_duerp", f.name, stdout=StringIO())
        self.assertEqual(ActionPrevention.objects.get(description="A1").categorie.nom, "Formation")
        self.assertTrue(CategorieAction.objects.filter(nom="Nouvelle").exists())


class ResponsableTests(Base):
    def test_responsable_texte_libre_et_suggestions(self):
        self.client.post(reverse("actions_ajouter"), {"portee": "risques", "risques": [self.rps_voirie.pk], "description": "Formation",
                                                      "statut": "a_faire", "responsable": "Service RH (M. Martin)"})
        self.assertEqual(ActionPrevention.objects.get(description="Formation").responsable, "Service RH (M. Martin)")
        ActionPrevention.objects.create(risque=self.rps_b, description="Action B", responsable="Pilote de B")
        self.client.force_login(self.user_a)
        page = self.client.get(reverse("actions_ajouter")).content.decode()
        self.assertIn('<option value="Service RH (M. Martin)">', page)
        self.assertNotIn("Pilote de B", page)   # pas de fuite des responsables d'une autre structure


class TableauDeBordTests(Base):
    def test_widgets_et_liens(self):
        import datetime
        from django.utils import timezone
        ActionPrevention.objects.create(risque=self.rps_accueil, description="Faite", statut="terminee", date_realisation=timezone.localdate(), cout=500)
        page = self.client.get(reverse("dashboard"))
        self.assertEqual(page.status_code, 200)
        for titre in ("Matrice gravité × fréquence", "Points de vigilance", "Avancement du plan d'actions", "Actions à compléter",
                      "Unités de travail les plus exposées", "Budget de prévention", "Mise à jour du DUERP et réévaluations"):
            self.assertContains(page, titre)
        ctx = page.context
        self.assertIn(self.rps_accueil, ctx["vigilance"])          # risque moyen dont la seule action est terminée
        self.assertIn(self.rps_accueil, ctx["reevaluer"])          # action terminée, risque jamais réévalué depuis
        self.assertEqual(ctx["budget"]["realise"], 500)
        # enregistrer la fiche du risque = réévaluation
        self.client.post(reverse("risques_modifier", args=[self.rps_accueil.pk]), {
            "unite": self.accueil.pk, "categorie": self.cat.pk, "danger": self.rps_accueil.danger, "frequence": 7, "gravite": 4, "maitrise": 4})
        self.rps_accueil.refresh_from_db()
        self.assertEqual(self.rps_accueil.date_evaluation, timezone.localdate())
        self.assertNotIn(self.rps_accueil, self.client.get(reverse("dashboard")).context["reevaluer"])

    def test_filtres_des_listes(self):
        ActionPrevention.objects.create(risque=self.rps_voirie, description="Sans échéance")
        self.assertContains(self.client.get(reverse("risques_liste") + "?gravite=4&frequence=7"), "Agressions accueil")
        self.assertNotContains(self.client.get(reverse("risques_liste") + "?gravite=4&frequence=7"), "Stress voirie")
        vigilance = self.client.get(reverse("risques_liste") + "?filtre=vigilance").content.decode()
        self.assertIn("Agressions accueil", vigilance)
        self.assertNotIn("Stress voirie", vigilance)   # a une action en cours
        self.assertContains(self.client.get(reverse("actions_liste") + "?manque=echeance"), "Sans échéance")


class IntroductionTests(Base):
    def test_saisie_document_et_archive(self):
        import base64
        png = b"\x89PNG\r\n\x1a\n" + b"0" * 50
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.post(reverse("introduction"), {"texte": "Présentation.\n\n## Références\n- Article L.4121-3",
                                                   "logo": SimpleUploadedFile("logo.png", png, content_type="image/png")})
        from core.models import Introduction
        intro = Introduction.objects.get(structure__isnull=True)
        self.assertTrue(intro.logo.startswith("data:image/png;base64,"))
        document = self.client.get(reverse("document"))
        self.assertContains(document, "page-garde")
        self.assertContains(document, "<li>Article L.4121-3</li>", html=True)
        self.assertContains(document, "Méthode d'évaluation")
        self.assertContains(document, 'href="#unite-1"')
        # un logo qui n'est pas une image est refusé
        r = self.client.post(reverse("introduction"), {"texte": "x", "logo": SimpleUploadedFile("logo.png", b"MZ\x90\x00", content_type="image/png")})
        self.assertContains(r, "PNG ou JPEG")
        # l'introduction est archivée avec la version
        self.client.post(reverse("versions_ajouter"), {"commentaire": "v1"})
        version = VersionDuerp.objects.get()
        self.assertIn("Présentation.", version.introduction["texte"])
        intro.texte = "Modifiée depuis"
        intro.save()
        self.assertContains(self.client.get(reverse("versions_document", args=[version.pk])), "Présentation.")

    def test_texte_echappe(self):
        from core.utils.introduction import mise_en_forme
        self.assertNotIn("<script>", mise_en_forme('<script>alert("x")</script>'))

    def test_import_introduction_seulement(self):
        donnees = {"format": "previthys-import/2", "unites": [], "categories": [], "risques": [], "actions": [],
                   "introduction": {"texte": "Texte importé", "logo": "logo.png"}}
        with tempfile.NamedTemporaryFile(suffix=".zip") as f:
            with zipfile.ZipFile(f.name, "w") as z:
                z.writestr("donnees.json", json.dumps(donnees))
                z.writestr("logo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 20)
            call_command("importer_duerp", f.name, "--introduction-seulement", stdout=StringIO())
        from core.models import Introduction
        self.assertEqual(Introduction.objects.get(structure__isnull=True).texte, "Texte importé")
        self.assertEqual(UniteTravail.objects.count(), 3)   # rien d'autre n'est importé


class FiltreUniteTests(Base):
    def test_filtre_unite_seul_et_combine(self):
        page = self.client.get(reverse("risques_liste") + "?unite=%d" % self.voirie.pk).content.decode()
        self.assertIn("Stress voirie", page)
        self.assertNotIn("Agressions accueil", page)
        self.assertIn("Voirie (1)", page)                       # nombre de risques par unité dans la liste déroulante
        combine = self.client.get(reverse("risques_liste") + "?filtre=vigilance&unite=%d" % self.accueil.pk).content.decode()
        self.assertIn("Agressions accueil", combine)
        self.assertIn('name="filtre" value="vigilance"', combine)   # l'autre filtre est conservé au changement d'unité

    def test_unite_d_une_autre_structure_absente(self):
        self.client.force_login(self.user_a)
        page = self.client.get(reverse("risques_liste") + "?unite=%d" % self.ecole_b.pk).content.decode()
        self.assertNotIn("École B", page)
        self.assertNotIn("Stress école B", page)


class FiltreUnitePlanTests(Base):
    def test_plan_filtre_par_unite(self):
        ActionPrevention.objects.create(risque=self.rps_voirie, description="Propre voirie")
        ActionPrevention.objects.create(risque=self.rps_accueil, description="Propre accueil")
        ActionPrevention.objects.create(risques=[self.rps_voirie, self.rps_accueil], description="Commune aux deux")
        ActionPrevention.objects.create(description="Générale pour tous")
        ActionPrevention.objects.create(description="Générale de B", structure=self.struct_b)
        page = self.client.get(reverse("actions_liste") + "?unite=%d" % self.voirie.pk).content.decode()
        for texte in ("Propre voirie", "Commune aux deux", "Générale pour tous"):
            self.assertIn(texte, page)
        for texte in ("Propre accueil", "Générale de B"):
            self.assertNotIn(texte, page)
        self.assertIn("Voirie (3)", page)
        # trois listes déroulantes : unité, portée, catégorie, avec les nombres tenant compte des autres filtres
        for nom in ("unite", "portee", "categorie"):
            self.assertIn('name="%s" data-soumettre' % nom, page)
        self.assertIn("Communes (1)", page)
        self.assertIn("Générales (1)", page)
        communes = self.client.get(reverse("actions_liste") + "?unite=%d&portee=communes" % self.voirie.pk).content.decode()
        self.assertIn("Commune aux deux", communes)
        self.assertNotIn("Propre voirie", communes)
        self.assertIn('<option value="communes" selected', communes)
        self.assertIn("Voirie (1)", communes)

    def test_filtre_a_completer_conserve_dans_les_listes(self):
        ActionPrevention.objects.create(risque=self.rps_voirie, description="Sans échéance")
        page = self.client.get(reverse("actions_liste") + "?manque=echeance&portee=unite").content.decode()
        self.assertIn('type="hidden" name="manque" value="echeance"', page)
        self.assertIn("?portee=unite", page)   # la pastille ✕ retire seulement « à compléter »
