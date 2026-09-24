#  -*- coding: utf-8 -*-
#  Copyright (c) 2026 Ivan LUCAS.
#  Previthys, application de gestion du DUERP (Document Unique d’Évaluation des Risques Professionnels).
#  Distribué sous licence GNU GPL.

import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

from core.models import SEUIL_CRITIQUE, SEUIL_MOYEN
from core.utils.donnees import normaliser_donnees

POLICE = "Arial"
BORD = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
NIVEAUX = {"Faible": ("C6EFCE", "006100"), "Moyen": ("FFEB9C", "9C5700"), "Critique": ("FFC7CE", "9C0006")}
NOTE_COTATION = "Cotation = fréquence × gravité × maîtrise (1 à 1 000). Faible jusqu'à 40, moyen de 41 à 196, critique à partir de 280."


class Formule(str):
    """Formule Excel écrite par l'application. Tout autre texte est écrit comme du texte, jamais comme une formule."""


def _texte(valeur):
    """Retire les caractères de contrôle interdits dans un classeur Excel (sinon l'export échouerait)."""
    return ILLEGAL_CHARACTERS_RE.sub("", valeur) if isinstance(valeur, str) else valeur


def _f(**kw):
    return Font(name=POLICE, size=kw.pop("size", 10), **kw)


def _entete(ws, ligne, titres):
    for col, titre in enumerate(titres, 1):
        c = ws.cell(row=ligne, column=col, value=titre)
        c.font = _f(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORD


def _cellule(ws, ligne, col, valeur, centre=False, fmt=None):
    c = ws.cell(row=ligne, column=col)
    if isinstance(valeur, Formule):
        c.value = str(valeur)
    elif isinstance(valeur, str):
        # Texte saisi par un utilisateur : jamais interprété comme une formule (=HYPERLINK, =cmd|...), même s'il commence par « = »
        c.value = _texte(valeur)
        c.data_type = "s"
    else:
        c.value = valeur
    c.font = _f()
    c.border = BORD
    c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if centre else None)
    if fmt:
        c.number_format = fmt
    return c


def _mesures(r):
    """Texte libre des mesures existantes, suivi des actions terminées (mesures réalisées)."""
    lignes = [r["mesures_existantes"]] if r.get("mesures_existantes") else []
    for m in r["mesures_realisees"]:
        lignes.append("• %s (réalisée le %s)" % (m["description"], m["date"]) if m["date"] else "• %s (réalisée)" % m["description"])
    return "\n".join(lignes)


def _date(texte):
    try:
        return datetime.datetime.strptime(texte, "%d/%m/%Y").date()
    except (TypeError, ValueError):
        return None


def _mise_en_forme_niveau(ws, plage):
    for libelle, (fond, texte) in NIVEAUX.items():
        ws.conditional_formatting.add(plage, CellIsRule(operator="equal", formula=['"%s"' % libelle], fill=PatternFill("solid", start_color=fond, end_color=fond), font=Font(name=POLICE, size=10, color=texte, bold=True)))


def generer_xlsx(donnees, sous_titre=""):
    """Construit le classeur Excel du DUERP à partir de la structure construire_donnees()."""
    risques, actions = [], []
    for unite in normaliser_donnees(donnees):
        for r in unite["risques"]:
            risques.append((unite["nom"], r))
            for a in r["actions_toutes"]:
                actions.append((unite["nom"], r["danger"], a))
    n_r = max(len(risques) + 1, 2)
    n_a = max(len(actions) + 1, 2)

    wb = Workbook()
    ws_s = wb.active
    ws_s.title = "Synthèse"
    ws_r = wb.create_sheet("Risques")
    ws_a = wb.create_sheet("Actions")

    # ----- Feuille Risques -----
    # Colonnes : A Unité, B Catégorie, C Danger, D Situation, E Fréquence, F Gravité, G Maîtrise, H Cotation (=E*F*G), I Niveau, J Mesures
    _entete(ws_r, 1, ["Unité de travail", "Catégorie", "Danger", "Situation de travail", "Fréquence", "Gravité", "Maîtrise", "Cotation (F × G × M)", "Niveau", "Mesures existantes"])
    for i, (unite, r) in enumerate(risques, 2):
        _cellule(ws_r, i, 1, unite)
        _cellule(ws_r, i, 2, r["categorie"])
        _cellule(ws_r, i, 3, r["danger"])
        _cellule(ws_r, i, 4, r["situation"])
        _cellule(ws_r, i, 5, r["frequence"], centre=True)
        _cellule(ws_r, i, 6, r["gravite"], centre=True)
        _cellule(ws_r, i, 7, r.get("maitrise", ""), centre=True)   # absent des archives créées avant l'ajout de ce facteur
        _cellule(ws_r, i, 8, Formule("=E%d*F%d*G%d" % (i, i, i)), centre=True)
        _cellule(ws_r, i, 9, Formule('=IF(H%d>=%d,"Critique",IF(H%d>=%d,"Moyen","Faible"))' % (i, SEUIL_CRITIQUE, i, SEUIL_MOYEN)), centre=True)
        _cellule(ws_r, i, 10, _mesures(r))
    for col, largeur in zip("ABCDEFGHIJ", (26, 24, 36, 40, 11, 11, 11, 14, 12, 40)):
        ws_r.column_dimensions[col].width = largeur
    ws_r.freeze_panes = "A2"
    ws_r.auto_filter.ref = "A1:J%d" % n_r
    _mise_en_forme_niveau(ws_r, "I2:I%d" % n_r)

    # ----- Feuille Actions -----
    _entete(ws_a, 1, ["Unité de travail", "Danger", "Action de prévention", "Responsable", "Échéance", "Statut", "En retard", "Date de réalisation"])
    for i, (unite, danger, a) in enumerate(actions, 2):
        _cellule(ws_a, i, 1, unite)
        _cellule(ws_a, i, 2, danger)
        _cellule(ws_a, i, 3, a["description"])
        _cellule(ws_a, i, 4, a.get("responsable", ""))
        _cellule(ws_a, i, 5, _date(a["echeance"]), centre=True, fmt="DD/MM/YYYY")
        _cellule(ws_a, i, 6, a["statut"], centre=True)
        _cellule(ws_a, i, 7, Formule('=IF(AND(F%d<>"Terminée",E%d<>"",E%d<TODAY()),"Oui","Non")' % (i, i, i)), centre=True)
        _cellule(ws_a, i, 8, _date(a.get("date_realisation", "")), centre=True, fmt="DD/MM/YYYY")
    for col, largeur in zip("ABCDEFGH", (26, 36, 46, 22, 14, 14, 12, 16)):
        ws_a.column_dimensions[col].width = largeur
    ws_a.freeze_panes = "A2"
    ws_a.auto_filter.ref = "A1:H%d" % n_a
    ws_a.conditional_formatting.add("G2:G%d" % n_a, CellIsRule(operator="equal", formula=['"Oui"'], fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE"), font=Font(name=POLICE, size=10, color="9C0006", bold=True)))

    # ----- Feuille Synthèse -----
    def section(ligne, titre):
        ws_s.cell(row=ligne, column=1, value=titre).font = _f(size=12, bold=True, color="1F3864")

    # Lignes fixes de la feuille (calculées à l'avance, pour que les indicateurs du haut puissent
    # référencer par formule le tableau des unités situé plus bas) :
    #   4  Indicateurs (titre)          10 Répartition par niveau (titre)   17 Unités de travail (titre)
    #   5-8  les 4 indicateurs          11 note sur la cotation             18 en-tête du tableau
    #                                   12 en-tête Niveau / Nombre          19.. une ligne par unité
    #                                   13-15 Faible / Moyen / Critique
    debut_unites = 19
    fin_unites = max(debut_unites + len(donnees) - 1, debut_unites)

    ws_s["A1"] = "Document unique d'évaluation des risques professionnels"
    ws_s["A1"].font = _f(size=16, bold=True)
    ws_s["A2"] = _texte(sous_titre)
    ws_s["A2"].data_type = "s"
    ws_s["A2"].font = _f(color="595959")

    section(4, "Indicateurs")
    indicateurs = [
        ("Unités de travail", Formule("=COUNTA(A%d:A%d)" % (debut_unites, fin_unites))),
        ("Risques évalués", Formule("=COUNTA(Risques!C2:C%d)" % n_r)),
        ("Risques critiques", Formule('=COUNTIF(Risques!I2:I%d,"Critique")' % n_r)),
        ("Actions en retard", Formule('=COUNTIF(Actions!G2:G%d,"Oui")' % n_a)),
    ]
    for i, (libelle, formule) in enumerate(indicateurs, 5):
        _cellule(ws_s, i, 1, libelle)
        _cellule(ws_s, i, 2, formule, centre=True).font = _f(bold=True)

    section(10, "Répartition des risques par niveau")
    ws_s["A11"] = NOTE_COTATION
    ws_s["A11"].font = _f(color="595959")
    _entete(ws_s, 12, ["Niveau", "Nombre de risques"])
    for k, niveau in enumerate(("Faible", "Moyen", "Critique")):
        ligne = 13 + k
        c1 = _cellule(ws_s, ligne, 1, niveau)
        c2 = _cellule(ws_s, ligne, 2, Formule('=COUNTIF(Risques!$I$2:$I$%d,"%s")' % (n_r, niveau)), centre=True)
        for c in (c1, c2):
            c.fill = PatternFill("solid", start_color=NIVEAUX[niveau][0])
            c.font = _f(bold=True, color=NIVEAUX[niveau][1])

    section(17, "Unités de travail")
    _entete(ws_s, 18, ["Unité de travail", "Effectif", "Risques", "Dont critiques"])
    for i, unite in enumerate(donnees, debut_unites):
        _cellule(ws_s, i, 1, unite["nom"])
        _cellule(ws_s, i, 2, unite["effectif"], centre=True)
        _cellule(ws_s, i, 3, Formule("=COUNTIF(Risques!$A$2:$A$%d,A%d)" % (n_r, i)), centre=True)
        _cellule(ws_s, i, 4, Formule('=COUNTIFS(Risques!$A$2:$A$%d,A%d,Risques!$I$2:$I$%d,"Critique")' % (n_r, i, n_r)), centre=True)
    ws_s.cell(row=fin_unites + 2, column=1, value=NOTE_COTATION).font = _f(color="595959")

    ws_s.column_dimensions["A"].width = 34
    for col in "BCD":
        ws_s.column_dimensions[col].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
